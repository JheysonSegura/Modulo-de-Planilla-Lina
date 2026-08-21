"""Asistente de migración de datos históricos (empresa que viene de otro
sistema). Ver PLAN-migracion-datos.txt para el diseño completo.

Alcance decidido: NO se replica el historial completo transacción por
transacción -- se migra (A) lo estructural exacto (empleado, contrato,
salario vigente), (B) saldos acumulados a una fecha de corte (vacaciones,
ISR del año), (C) una ventana acotada de bruto histórico (hasta 5 años,
opcional). Todos los cálculos del motor (décimo, liquidaciones, ISR) ya
tienen ventanas acotadas o resguardos si falta un dato -- ver el plan.

Solo aplica a empresas SIN datos operativos todavía (empresas_vacia):
es una herramienta de onboarding de una sola vez.
"""
import calendar
import dataclasses
import datetime
import decimal
import io
import uuid

import pydantic
from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.models import (
    Contrato,
    Empleado,
    Empresa,
    HistorialSalarial,
    MovimientoPlanilla,
    Planilla,
    ProvisionVacaciones,
)
from app.repositories import contratos as contratos_repo
from app.repositories import empleados as empleados_repo
from app.repositories import empresas as empresas_repo
from app.repositories import historial_salarial as historial_repo
from app.repositories import movimientos_planilla as movimientos_repo
from app.repositories import planillas as planillas_repo
from app.schemas.contratos import ContratoCreate
from app.schemas.empleados import EmpleadoCreate
from app.schemas.migracion import (
    ErrorValidacionMigracion,
    ResultadoValidacionMigracion,
    ResumenMigracion,
)
from app.services import auditoria_service
from app.services.salario_minimo_service import validar_salario_minimo

_CERO = decimal.Decimal("0")
_CENTAVO = decimal.Decimal("0.01")
DIAS_MES_COMERCIAL = decimal.Decimal("30")

HOJA_INSTRUCCIONES = "Instrucciones"
HOJA_EMPLEADOS = "Empleados y contratos"
HOJA_CAMBIOS_SALARIO = "Cambios de salario recientes"
HOJA_VACACIONES = "Saldo de vacaciones"
HOJA_ISR = "ISR retenido este año"
HOJA_BRUTO_HISTORICO = "Bruto histórico mensual"

_VERDADEROS = {"si", "sí", "true", "1", "x"}


# --- Definición de columnas (mismo orden en la plantilla y en la lectura) --


COLUMNAS_EMPLEADOS = [
    ("identificacion", "Identificación *"),
    ("tipo_identificacion", "Tipo identificación (cedula/pasaporte)"),
    ("nombre_completo", "Nombre completo *"),
    ("fecha_nacimiento", "Fecha de nacimiento (AAAA-MM-DD)"),
    ("nacionalidad", "Nacionalidad"),
    ("sexo", "Sexo (masculino/femenino/otro)"),
    ("telefono", "Teléfono"),
    ("cargo", "Cargo *"),
    ("tipo_contrato", "Tipo de contrato (indefinido/definido/obra_determinada) *"),
    ("fecha_inicio", "Fecha de inicio REAL del contrato (AAAA-MM-DD) *"),
    ("salario_base", "Salario mensual VIGENTE HOY *"),
    ("jornada_horas_semana", "Jornada horas/semana (default 48)"),
    ("periodicidad_pago", "Periodicidad de pago (quincenal/mensual, default quincenal)"),
    ("es_tecnico", "¿Es técnico? Art. 222 CT (SI/NO)"),
    ("exento_salario_minimo", "¿Exento de salario mínimo? (SI/NO)"),
    ("motivo_exencion_salario_minimo", "Motivo de exención (obligatorio si el anterior es SI)"),
]

COLUMNAS_CAMBIOS_SALARIO = [
    ("identificacion", "Identificación (debe existir en la hoja Empleados y contratos) *"),
    ("salario_anterior", "Salario de ese tramo *"),
    ("fecha_vigencia_desde", "Vigente desde (AAAA-MM-DD) *"),
    ("fecha_vigencia_hasta", "Vigente hasta (AAAA-MM-DD) *"),
]

COLUMNAS_VACACIONES = [
    ("identificacion", "Identificación (debe existir en la hoja Empleados y contratos) *"),
    ("dias_acumulados", "Días acumulados a la fecha de corte *"),
    ("dias_gozados", "Días ya gozados de esos acumulados *"),
]

COLUMNAS_ISR = [
    ("identificacion", "Identificación (debe existir en la hoja Empleados y contratos) *"),
    ("monto_isr_retenido", "ISR ya retenido en el año calendario en curso *"),
]

COLUMNAS_BRUTO_HISTORICO = [
    ("identificacion", "Identificación (debe existir en la hoja Empleados y contratos) *"),
    ("mes", "Mes (cualquier fecha dentro del mes, AAAA-MM-DD) *"),
    ("salario_bruto", "Salario bruto de ese mes *"),
]

_TEXTO_INSTRUCCIONES = [
    "Plantilla de migración de datos -- Nómina Panamá",
    "",
    "Hoja 'Empleados y contratos': OBLIGATORIA. Un empleado por fila. El",
    "salario_base es el que rige HOY, no el histórico.",
    "",
    "Hoja 'Cambios de salario recientes': OPCIONAL. Solo si el salario",
    "actual no aplicó durante los últimos 4 meses (afecta la precisión",
    "del próximo pago de décimo). Si se deja vacía, se asume que el",
    "salario actual rigió todo ese período.",
    "",
    "Hoja 'Saldo de vacaciones': OPCIONAL pero MUY recomendada. Si se deja",
    "vacía para un empleado, su saldo de vacaciones acumulado en el",
    "sistema anterior se pierde -- empieza en cero desde la fecha de corte.",
    "",
    "Hoja 'ISR retenido este año': OPCIONAL. Si se deja vacía, el sistema",
    "asume que no se ha retenido nada este año calendario, lo que puede",
    "sub-retener ISR el resto del año para ese empleado.",
    "",
    "Hoja 'Bruto histórico mensual': OPCIONAL, hasta 5 años atrás. Solo",
    "importa si se va a liquidar (despedir/renuncia) a alguien pronto --",
    "sin estos datos, las liquidaciones usan el salario actual como",
    "aproximación (nunca calculan cero ni fallan).",
    "",
    "Los campos marcados con * son obligatorios en su hoja (si la hoja",
    "tiene al menos una fila con datos).",
]


# --- Guard: solo empresas nuevas sin datos operativos ---------------------


def verificar_empresa_disponible(db: Session, empresa_id: uuid.UUID) -> tuple[bool, str | None]:
    if empleados_repo.listar(db):
        return False, "Esta empresa ya tiene empleados registrados -- el asistente de migración es solo para el arranque."
    if planillas_repo.listar(db, empresa_id):
        return False, "Esta empresa ya tiene planillas registradas -- el asistente de migración es solo para el arranque."
    return True, None


def _exigir_empresa_disponible(db: Session, empresa_id: uuid.UUID) -> None:
    disponible, motivo = verificar_empresa_disponible(db, empresa_id)
    if not disponible:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, motivo)


# --- Generación de la plantilla descargable --------------------------------


def generar_plantilla() -> bytes:
    workbook = Workbook()
    hoja_instrucciones = workbook.active
    hoja_instrucciones.title = HOJA_INSTRUCCIONES
    for linea in _TEXTO_INSTRUCCIONES:
        hoja_instrucciones.append([linea])
    hoja_instrucciones.column_dimensions["A"].width = 90

    for titulo_hoja, columnas in (
        (HOJA_EMPLEADOS, COLUMNAS_EMPLEADOS),
        (HOJA_CAMBIOS_SALARIO, COLUMNAS_CAMBIOS_SALARIO),
        (HOJA_VACACIONES, COLUMNAS_VACACIONES),
        (HOJA_ISR, COLUMNAS_ISR),
        (HOJA_BRUTO_HISTORICO, COLUMNAS_BRUTO_HISTORICO),
    ):
        hoja = workbook.create_sheet(titulo_hoja)
        hoja.append([titulo for _, titulo in columnas])
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for i in range(len(columnas)):
            hoja.column_dimensions[hoja.cell(row=1, column=i + 1).column_letter].width = 32

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Helpers de parseo -----------------------------------------------------


def _fila_vacia(valores: tuple) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in valores)


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _booleano(valor) -> bool:
    texto = _texto(valor)
    return texto is not None and texto.lower() in _VERDADEROS


def _fecha(valor) -> datetime.date | None:
    if valor is None:
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    texto = _texto(valor)
    if texto is None:
        return None
    try:
        return datetime.date.fromisoformat(texto)
    except ValueError:
        return None


def _decimal(valor) -> decimal.Decimal | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        return decimal.Decimal(str(valor))
    except (decimal.InvalidOperation, ValueError):
        return None


def _ultimo_dia_del_mes(fecha: datetime.date) -> datetime.date:
    ultimo_dia = calendar.monthrange(fecha.year, fecha.month)[1]
    return fecha.replace(day=ultimo_dia)


def _primer_dia_del_mes(fecha: datetime.date) -> datetime.date:
    return fecha.replace(day=1)


def _leer_filas(workbook, nombre_hoja: str, num_columnas: int):
    if nombre_hoja not in workbook.sheetnames:
        return
    hoja = workbook[nombre_hoja]
    for i, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        valores = tuple(fila[:num_columnas]) if fila else ()
        if not valores or _fila_vacia(valores):
            continue
        yield i, valores


# --- Filas parseadas (usadas tanto en validación como en la escritura) ----


@dataclasses.dataclass
class FilaEmpleado:
    fila: int
    empleado: EmpleadoCreate
    contrato: ContratoCreate


@dataclasses.dataclass
class FilaCambioSalario:
    fila: int
    identificacion_clave: tuple[str, str]
    salario_anterior: decimal.Decimal
    fecha_vigencia_desde: datetime.date
    fecha_vigencia_hasta: datetime.date


@dataclasses.dataclass
class FilaVacaciones:
    fila: int
    identificacion_clave: tuple[str, str]
    dias_acumulados: decimal.Decimal
    dias_gozados: decimal.Decimal


@dataclasses.dataclass
class FilaIsr:
    fila: int
    identificacion_clave: tuple[str, str]
    monto: decimal.Decimal


@dataclasses.dataclass
class FilaBrutoHistorico:
    fila: int
    identificacion_clave: tuple[str, str]
    mes: datetime.date
    salario_bruto: decimal.Decimal


@dataclasses.dataclass
class _ResultadoLectura:
    errores: list[ErrorValidacionMigracion]
    filas_empleados: list[FilaEmpleado]
    filas_cambios_salario: list[FilaCambioSalario]
    filas_vacaciones: list[FilaVacaciones]
    filas_isr: list[FilaIsr]
    filas_bruto_historico: list[FilaBrutoHistorico]


def _error(errores: list, hoja: str, fila: int, campo: str | None, mensaje: str) -> None:
    errores.append(ErrorValidacionMigracion(hoja=hoja, fila=fila, campo=campo, mensaje=mensaje))


# --- Lectura + validación (dry-run, no escribe nada) -----------------------


def _leer_y_validar(
    db: Session, empresa: Empresa, archivo_bytes: bytes, fecha_corte: datetime.date
) -> _ResultadoLectura:
    errores: list[ErrorValidacionMigracion] = []
    try:
        workbook = load_workbook(io.BytesIO(archivo_bytes), data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY, "El archivo no es un Excel (.xlsx) válido."
        ) from exc

    if HOJA_EMPLEADOS not in workbook.sheetnames:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            f"Falta la hoja obligatoria '{HOJA_EMPLEADOS}'.",
        )

    filas_empleados: list[FilaEmpleado] = []
    claves_vistas: dict[tuple[str, str], int] = {}

    for i, v in _leer_filas(workbook, HOJA_EMPLEADOS, len(COLUMNAS_EMPLEADOS)):
        (
            identificacion, tipo_identificacion, nombre_completo, fecha_nacimiento,
            nacionalidad, sexo, telefono, cargo, tipo_contrato, fecha_inicio,
            salario_base, jornada, periodicidad, es_tecnico, exento, motivo_exencion,
        ) = v

        datos_empleado = {
            "identificacion": _texto(identificacion),
            "tipo_identificacion": _texto(tipo_identificacion) or "cedula",
            "nombre_completo": _texto(nombre_completo),
            "fecha_nacimiento": _fecha(fecha_nacimiento),
            "nacionalidad": _texto(nacionalidad),
            "sexo": _texto(sexo),
            "telefono": _texto(telefono),
        }
        datos_contrato = {
            "tipo_contrato": _texto(tipo_contrato),
            "cargo": _texto(cargo),
            "fecha_inicio": _fecha(fecha_inicio),
            "salario_base": _decimal(salario_base),
            "jornada_horas_semana": _decimal(jornada) or decimal.Decimal("48"),
            "periodicidad_pago": _texto(periodicidad) or "quincenal",
            "es_tecnico": _booleano(es_tecnico),
            "exento_salario_minimo": _booleano(exento),
            "motivo_exencion_salario_minimo": _texto(motivo_exencion),
        }

        empleado_obj: EmpleadoCreate | None = None
        try:
            empleado_obj = EmpleadoCreate(**datos_empleado)
        except pydantic.ValidationError as exc:
            for err in exc.errors():
                campo = ".".join(str(p) for p in err["loc"])
                _error(errores, HOJA_EMPLEADOS, i, campo, err["msg"])

        contrato_obj: ContratoCreate | None = None
        try:
            contrato_obj = ContratoCreate(**datos_contrato)
        except pydantic.ValidationError as exc:
            for err in exc.errors():
                campo = ".".join(str(p) for p in err["loc"])
                _error(errores, HOJA_EMPLEADOS, i, campo, err["msg"])

        if empleado_obj is None or contrato_obj is None:
            continue

        clave = (empleado_obj.tipo_identificacion, empleado_obj.identificacion)
        if clave in claves_vistas:
            _error(
                errores, HOJA_EMPLEADOS, i, "identificacion",
                f"Identificación duplicada en el archivo (ya aparece en la fila {claves_vistas[clave]}).",
            )
            continue
        claves_vistas[clave] = i

        if contrato_obj.fecha_inicio >= fecha_corte:
            _error(
                errores, HOJA_EMPLEADOS, i, "fecha_inicio",
                "fecha_inicio debe ser anterior a la fecha de corte de la migración.",
            )
            continue

        if not contrato_obj.exento_salario_minimo:
            try:
                validar_salario_minimo(
                    db, empresa, contrato_obj.fecha_inicio, contrato_obj.salario_base,
                    contrato_obj.jornada_horas_semana,
                )
            except HTTPException as exc:
                _error(errores, HOJA_EMPLEADOS, i, "salario_base", str(exc.detail))
                continue

        filas_empleados.append(FilaEmpleado(fila=i, empleado=empleado_obj, contrato=contrato_obj))

    # --- Hoja 2: cambios de salario recientes (opcional) -------------------
    filas_cambios_salario: list[FilaCambioSalario] = []
    for i, v in _leer_filas(workbook, HOJA_CAMBIOS_SALARIO, len(COLUMNAS_CAMBIOS_SALARIO)):
        identificacion, salario_anterior_v, desde_v, hasta_v = v
        clave = _resolver_clave(identificacion, claves_vistas, errores, HOJA_CAMBIOS_SALARIO, i)
        salario_anterior = _decimal(salario_anterior_v)
        desde = _fecha(desde_v)
        hasta = _fecha(hasta_v)

        if salario_anterior is None or salario_anterior <= _CERO:
            _error(errores, HOJA_CAMBIOS_SALARIO, i, "salario_anterior", "Debe ser un monto mayor a 0.")
        if desde is None:
            _error(errores, HOJA_CAMBIOS_SALARIO, i, "fecha_vigencia_desde", "Fecha inválida.")
        if hasta is None:
            _error(errores, HOJA_CAMBIOS_SALARIO, i, "fecha_vigencia_hasta", "Fecha inválida.")
        if desde is not None and hasta is not None and hasta < desde:
            _error(errores, HOJA_CAMBIOS_SALARIO, i, "fecha_vigencia_hasta", "Debe ser posterior a fecha_vigencia_desde.")
        if hasta is not None and hasta >= fecha_corte:
            _error(errores, HOJA_CAMBIOS_SALARIO, i, "fecha_vigencia_hasta", "Debe ser anterior a la fecha de corte.")

        if clave and salario_anterior is not None and desde is not None and hasta is not None and hasta >= desde:
            filas_cambios_salario.append(
                FilaCambioSalario(
                    fila=i, identificacion_clave=clave, salario_anterior=salario_anterior,
                    fecha_vigencia_desde=desde, fecha_vigencia_hasta=hasta,
                )
            )

    # --- Hoja 3: saldo de vacaciones (opcional) -----------------------------
    filas_vacaciones: list[FilaVacaciones] = []
    for i, v in _leer_filas(workbook, HOJA_VACACIONES, len(COLUMNAS_VACACIONES)):
        identificacion, dias_acumulados_v, dias_gozados_v = v
        clave = _resolver_clave(identificacion, claves_vistas, errores, HOJA_VACACIONES, i)
        dias_acumulados = _decimal(dias_acumulados_v)
        dias_gozados = _decimal(dias_gozados_v) or _CERO

        if dias_acumulados is None or dias_acumulados < _CERO:
            _error(errores, HOJA_VACACIONES, i, "dias_acumulados", "Debe ser un número mayor o igual a 0.")
        elif dias_gozados < _CERO:
            _error(errores, HOJA_VACACIONES, i, "dias_gozados", "Debe ser un número mayor o igual a 0.")
        elif dias_gozados > dias_acumulados:
            _error(errores, HOJA_VACACIONES, i, "dias_gozados", "No puede ser mayor a dias_acumulados.")
        elif clave:
            filas_vacaciones.append(
                FilaVacaciones(fila=i, identificacion_clave=clave, dias_acumulados=dias_acumulados, dias_gozados=dias_gozados)
            )

    # --- Hoja 4: ISR retenido este año (opcional) ---------------------------
    filas_isr: list[FilaIsr] = []
    for i, v in _leer_filas(workbook, HOJA_ISR, len(COLUMNAS_ISR)):
        identificacion, monto_v = v
        clave = _resolver_clave(identificacion, claves_vistas, errores, HOJA_ISR, i)
        monto = _decimal(monto_v)
        if monto is None or monto < _CERO:
            _error(errores, HOJA_ISR, i, "monto_isr_retenido", "Debe ser un monto mayor o igual a 0.")
        elif clave:
            filas_isr.append(FilaIsr(fila=i, identificacion_clave=clave, monto=monto))

    # --- Hoja 5: bruto histórico mensual (opcional) -------------------------
    filas_bruto_historico: list[FilaBrutoHistorico] = []
    for i, v in _leer_filas(workbook, HOJA_BRUTO_HISTORICO, len(COLUMNAS_BRUTO_HISTORICO)):
        identificacion, mes_v, salario_bruto_v = v
        clave = _resolver_clave(identificacion, claves_vistas, errores, HOJA_BRUTO_HISTORICO, i)
        mes = _fecha(mes_v)
        salario_bruto = _decimal(salario_bruto_v)

        if mes is None:
            _error(errores, HOJA_BRUTO_HISTORICO, i, "mes", "Fecha inválida.")
        elif mes >= fecha_corte:
            _error(errores, HOJA_BRUTO_HISTORICO, i, "mes", "Debe ser anterior a la fecha de corte.")
        if salario_bruto is None or salario_bruto < _CERO:
            _error(errores, HOJA_BRUTO_HISTORICO, i, "salario_bruto", "Debe ser un monto mayor o igual a 0.")

        if clave and mes is not None and mes < fecha_corte and salario_bruto is not None and salario_bruto >= _CERO:
            filas_bruto_historico.append(
                FilaBrutoHistorico(fila=i, identificacion_clave=clave, mes=mes, salario_bruto=salario_bruto)
            )

    return _ResultadoLectura(
        errores=errores,
        filas_empleados=filas_empleados,
        filas_cambios_salario=filas_cambios_salario,
        filas_vacaciones=filas_vacaciones,
        filas_isr=filas_isr,
        filas_bruto_historico=filas_bruto_historico,
    )


def _resolver_clave(
    identificacion_raw, claves_vistas: dict, errores: list, hoja: str, fila: int
) -> tuple[str, str] | None:
    identificacion = _texto(identificacion_raw)
    if identificacion is None:
        _error(errores, hoja, fila, "identificacion", "Identificación obligatoria.")
        return None
    for clave in claves_vistas:
        if clave[1] == identificacion:
            return clave
    _error(
        errores, hoja, fila, "identificacion",
        f"'{identificacion}' no existe en la hoja '{HOJA_EMPLEADOS}'.",
    )
    return None


def _armar_resumen(resultado: _ResultadoLectura) -> ResumenMigracion:
    dias_vacaciones_totales = sum(
        (f.dias_acumulados - f.dias_gozados for f in resultado.filas_vacaciones), _CERO
    )
    isr_total = sum((f.monto for f in resultado.filas_isr), _CERO)
    return ResumenMigracion(
        empleados=len(resultado.filas_empleados),
        contratos=len(resultado.filas_empleados),
        tramos_salario_adicionales=len(resultado.filas_cambios_salario),
        dias_vacaciones_acumulados_totales=str(dias_vacaciones_totales.quantize(_CENTAVO)),
        isr_total_a_cargar=str(isr_total.quantize(_CENTAVO)),
        meses_bruto_historico_cargados=len(resultado.filas_bruto_historico),
    )


def leer_y_validar(
    db: Session, empresa_id: uuid.UUID, archivo_bytes: bytes, fecha_corte: datetime.date
) -> ResultadoValidacionMigracion:
    """Dry-run público (usado por POST /migracion/validar): nunca escribe
    en la base de datos."""
    _exigir_empresa_disponible(db, empresa_id)
    empresa = empresas_repo.get(db, empresa_id)
    resultado = _leer_y_validar(db, empresa, archivo_bytes, fecha_corte)
    resumen = _armar_resumen(resultado) if not resultado.errores else None
    return ResultadoValidacionMigracion(errores=resultado.errores, resumen=resumen)


# --- Ejecución real (escribe todo en una sola transacción) ----------------


def ejecutar_migracion(
    db: Session,
    empresa_id: uuid.UUID,
    usuario_id: uuid.UUID,
    archivo_bytes: bytes,
    fecha_corte: datetime.date,
) -> ResumenMigracion:
    _exigir_empresa_disponible(db, empresa_id)
    empresa = empresas_repo.get(db, empresa_id)
    resultado = _leer_y_validar(db, empresa, archivo_bytes, fecha_corte)
    if resultado.errores:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "El archivo tiene errores de validación pendientes -- volvé a llamar a "
            "/migracion/validar antes de confirmar.",
        )

    contratos_por_clave: dict[tuple[str, str], Contrato] = {}
    salario_actual_por_clave: dict[tuple[str, str], decimal.Decimal] = {}

    for fila_e in resultado.filas_empleados:
        clave = (fila_e.empleado.tipo_identificacion, fila_e.empleado.identificacion)

        empleado = Empleado(empresa_id=empresa_id, **fila_e.empleado.model_dump())
        empleados_repo.crear(db, empleado)

        contrato = Contrato(
            empresa_id=empresa_id, empleado_id=empleado.id,
            **fila_e.contrato.model_dump(exclude={"salario_base"}),
        )
        contratos_repo.crear(db, contrato)
        contratos_por_clave[clave] = contrato
        salario_actual_por_clave[clave] = fila_e.contrato.salario_base

        historial_repo.crear(
            db,
            HistorialSalarial(
                contrato_id=contrato.id, salario_base=fila_e.contrato.salario_base,
                fecha_vigencia_desde=fila_e.contrato.fecha_inicio, fecha_vigencia_hasta=None,
                motivo="migracion (vigente)",
            ),
        )

        # Ancla de seguridad -- SIEMPRE se crea, tenga o no datos de la
        # hoja ISR: sin esto, liquidaciones_service._calcular_salario_pendiente
        # caería a contrato.fecha_inicio (años atrás) si el contrato se
        # liquida antes de correr una sola planilla real en el sistema
        # nuevo, contando de más "salario pendiente" de forma severa.
        fin_ancla = fecha_corte - datetime.timedelta(days=1)
        planilla_ancla = Planilla(
            empresa_id=empresa_id, tipo="mensual",
            periodo_inicio=_primer_dia_del_mes(fin_ancla), periodo_fin=fin_ancla,
            fecha_pago=fin_ancla, estado="pagada",
        )
        planillas_repo.crear(db, planilla_ancla)
        movimientos_repo.crear(
            db,
            MovimientoPlanilla(
                planilla_id=planilla_ancla.id, contrato_id=contrato.id,
                salario_base_periodo=_CERO, salario_bruto=_CERO, css_empleado=_CERO,
                css_patronal=_CERO, seguro_educativo_empleado=_CERO,
                seguro_educativo_patronal=_CERO, riesgo_profesional_patronal=_CERO,
                isr_retenido=_CERO, otras_deducciones=_CERO, salario_neto=_CERO,
            ),
        )

        # Ancla de seguridad para vacaciones -- SIEMPRE se crea el período
        # 'abierto' aquí mismo: si no existiera, la primera planilla real
        # lo crearía sola con fecha_inicio_periodo=contrato.fecha_inicio
        # (años atrás), y _calcular_dias_y_monto_acumulado recorrería
        # TODO el historial_salarial ya cargado, acumulando de más.
        provision_abierta = ProvisionVacaciones(
            empresa_id=empresa_id, contrato_id=contrato.id, fecha_inicio_periodo=fecha_corte,
            dias_acumulados=_CERO, dias_gozados=_CERO, monto_provisionado=_CERO, estado="abierto",
        )
        db.add(provision_abierta)

    # --- Hoja 2: cambios de salario recientes ------------------------------
    por_empleado: dict[tuple[str, str], list[FilaCambioSalario]] = {}
    for fila_c in resultado.filas_cambios_salario:
        por_empleado.setdefault(fila_c.identificacion_clave, []).append(fila_c)

    for clave, filas in por_empleado.items():
        contrato = contratos_por_clave[clave]
        filas_ordenadas = sorted(filas, key=lambda f: f.fecha_vigencia_desde)
        for fila_c in filas_ordenadas:
            historial_repo.crear(
                db,
                HistorialSalarial(
                    contrato_id=contrato.id, salario_base=fila_c.salario_anterior,
                    fecha_vigencia_desde=fila_c.fecha_vigencia_desde,
                    fecha_vigencia_hasta=fila_c.fecha_vigencia_hasta, motivo="migracion (tramo anterior)",
                ),
            )
        ultimo_tramo_hasta = max(f.fecha_vigencia_hasta for f in filas_ordenadas)
        tramo_abierto = historial_repo.get_abierto(db, contrato.id)
        if tramo_abierto is not None:
            tramo_abierto.fecha_vigencia_desde = ultimo_tramo_hasta + datetime.timedelta(days=1)

    # --- Hoja 3: saldo de vacaciones ----------------------------------------
    for fila_v in resultado.filas_vacaciones:
        contrato = contratos_por_clave[fila_v.identificacion_clave]
        valor_dia = salario_actual_por_clave[fila_v.identificacion_clave] / DIAS_MES_COMERCIAL
        saldo = fila_v.dias_acumulados - fila_v.dias_gozados
        monto_provisionado = (saldo * valor_dia).quantize(_CENTAVO) if saldo > _CERO else _CERO
        db.add(
            ProvisionVacaciones(
                empresa_id=empresa_id, contrato_id=contrato.id,
                fecha_inicio_periodo=contrato.fecha_inicio, dias_acumulados=fila_v.dias_acumulados,
                dias_gozados=fila_v.dias_gozados, monto_provisionado=monto_provisionado, estado="acumulado",
            )
        )

    # --- Hoja 4: ISR retenido este año --------------------------------------
    for fila_i in resultado.filas_isr:
        contrato = contratos_por_clave[fila_i.identificacion_clave]
        planilla_isr = Planilla(
            empresa_id=empresa_id, tipo="mensual",
            periodo_inicio=datetime.date(fecha_corte.year, 1, 1),
            periodo_fin=fecha_corte - datetime.timedelta(days=1),
            fecha_pago=fecha_corte - datetime.timedelta(days=1), estado="pagada",
        )
        planillas_repo.crear(db, planilla_isr)
        movimientos_repo.crear(
            db,
            MovimientoPlanilla(
                planilla_id=planilla_isr.id, contrato_id=contrato.id,
                salario_base_periodo=_CERO, salario_bruto=_CERO, css_empleado=_CERO,
                css_patronal=_CERO, seguro_educativo_empleado=_CERO,
                seguro_educativo_patronal=_CERO, riesgo_profesional_patronal=_CERO,
                isr_retenido=fila_i.monto, otras_deducciones=_CERO, salario_neto=_CERO,
            ),
        )

    # --- Hoja 5: bruto histórico mensual ------------------------------------
    for fila_b in resultado.filas_bruto_historico:
        contrato = contratos_por_clave[fila_b.identificacion_clave]
        inicio_mes = _primer_dia_del_mes(fila_b.mes)
        fin_mes = _ultimo_dia_del_mes(fila_b.mes)
        planilla_mes = Planilla(
            empresa_id=empresa_id, tipo="mensual", periodo_inicio=inicio_mes,
            periodo_fin=fin_mes, fecha_pago=fin_mes, estado="pagada",
        )
        planillas_repo.crear(db, planilla_mes)
        movimientos_repo.crear(
            db,
            MovimientoPlanilla(
                planilla_id=planilla_mes.id, contrato_id=contrato.id,
                salario_base_periodo=fila_b.salario_bruto, salario_bruto=fila_b.salario_bruto,
                css_empleado=_CERO, css_patronal=_CERO, seguro_educativo_empleado=_CERO,
                seguro_educativo_patronal=_CERO, riesgo_profesional_patronal=_CERO,
                isr_retenido=_CERO, otras_deducciones=_CERO, salario_neto=fila_b.salario_bruto,
            ),
        )

    resumen = _armar_resumen(resultado)
    auditoria_service.registrar(
        db, empresa_id, usuario_id, "empresas", empresa_id, "migracion_confirmada",
        datos_nuevos={"fecha_corte": fecha_corte.isoformat(), **resumen.model_dump()},
    )
    db.commit()
    # Sin db.refresh(): rompería RLS igual que en el resto del proyecto.
    return resumen
