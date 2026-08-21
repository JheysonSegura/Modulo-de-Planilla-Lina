import datetime
import decimal
import io
import uuid

from openpyxl import Workbook

from app.services import migracion_service
from tests.conftest import crear_empresa, crear_salario_minimo, crear_usuario, login_y_seleccionar, vincular


def _sufijo() -> str:
    return uuid.uuid4().hex[:8]


def _preparar_empresa_admin(db, client) -> dict:
    crear_salario_minimo(db, datetime.date(2000, 1, 1), decimal.Decimal("1.00"), fecha_fin=None)
    empresa = crear_empresa(db, _sufijo())
    usuario = crear_usuario(db, f"admin-{_sufijo()}@example.com")
    vincular(db, usuario, empresa, "admin")
    return login_y_seleccionar(client, usuario, empresa)


def _construir_excel(
    filas_empleados: list[dict],
    filas_cambio_salario: list[dict] | None = None,
    filas_vacaciones: list[dict] | None = None,
    filas_isr: list[dict] | None = None,
    filas_bruto: list[dict] | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    hoja_empleados = workbook.create_sheet(migracion_service.HOJA_EMPLEADOS)
    hoja_empleados.append([c for c, _ in migracion_service.COLUMNAS_EMPLEADOS])
    for fila in filas_empleados:
        hoja_empleados.append([fila.get(c) for c, _ in migracion_service.COLUMNAS_EMPLEADOS])

    hoja_cambios = workbook.create_sheet(migracion_service.HOJA_CAMBIOS_SALARIO)
    hoja_cambios.append([c for c, _ in migracion_service.COLUMNAS_CAMBIOS_SALARIO])
    for fila in filas_cambio_salario or []:
        hoja_cambios.append([fila.get(c) for c, _ in migracion_service.COLUMNAS_CAMBIOS_SALARIO])

    hoja_vac = workbook.create_sheet(migracion_service.HOJA_VACACIONES)
    hoja_vac.append([c for c, _ in migracion_service.COLUMNAS_VACACIONES])
    for fila in filas_vacaciones or []:
        hoja_vac.append([fila.get(c) for c, _ in migracion_service.COLUMNAS_VACACIONES])

    hoja_isr = workbook.create_sheet(migracion_service.HOJA_ISR)
    hoja_isr.append([c for c, _ in migracion_service.COLUMNAS_ISR])
    for fila in filas_isr or []:
        hoja_isr.append([fila.get(c) for c, _ in migracion_service.COLUMNAS_ISR])

    hoja_bruto = workbook.create_sheet(migracion_service.HOJA_BRUTO_HISTORICO)
    hoja_bruto.append([c for c, _ in migracion_service.COLUMNAS_BRUTO_HISTORICO])
    for fila in filas_bruto or []:
        hoja_bruto.append([fila.get(c) for c, _ in migracion_service.COLUMNAS_BRUTO_HISTORICO])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _empleado_fila(identificacion: str, fecha_inicio: str, salario_base: float, **overrides) -> dict:
    fila = {
        "identificacion": identificacion,
        "tipo_identificacion": "cedula",
        "nombre_completo": f"Empleado {identificacion}",
        "cargo": "Contador",
        "tipo_contrato": "indefinido",
        "fecha_inicio": fecha_inicio,
        "salario_base": salario_base,
        "jornada_horas_semana": 48,
        "periodicidad_pago": "quincenal",
        "es_tecnico": "NO",
        "exento_salario_minimo": "NO",
    }
    fila.update(overrides)
    return fila


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _subir(client, headers, path, contenido, fecha_corte):
    return client.post(
        path,
        headers=headers,
        data={"fecha_corte": fecha_corte.isoformat()},
        files={"archivo": ("plantilla.xlsx", contenido, _XLSX)},
    )


def test_disponible_empresa_nueva(client, db):
    headers = _preparar_empresa_admin(db, client)
    resp = client.get("/empresas/actual/migracion/disponible", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["disponible"] is True


def test_descargar_plantilla(client, db):
    headers = _preparar_empresa_admin(db, client)
    resp = client.get("/empresas/actual/migracion/plantilla", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_migracion_completa_camino_feliz(client, db):
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"

    contenido = _construir_excel(
        filas_empleados=[_empleado_fila(identificacion, "2010-03-15", 900.00)],
        filas_vacaciones=[{"identificacion": identificacion, "dias_acumulados": 40, "dias_gozados": 10}],
        filas_isr=[{"identificacion": identificacion, "monto_isr_retenido": 120.50}],
        filas_bruto=[
            {"identificacion": identificacion, "mes": "2026-07-15", "salario_bruto": 900.00},
            {"identificacion": identificacion, "mes": "2026-08-15", "salario_bruto": 900.00},
        ],
    )

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    resultado = resp.json()
    assert resultado["errores"] == []
    assert resultado["resumen"]["empleados"] == 1
    assert resultado["resumen"]["meses_bruto_historico_cargados"] == 2

    resp = _subir(client, headers, "/empresas/actual/migracion/confirmar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    assert resp.json()["resumen"]["empleados"] == 1

    resp = client.get("/empleados", headers=headers)
    assert resp.status_code == 200, resp.text
    assert any(e["identificacion"] == identificacion for e in resp.json())

    resp = client.get("/empresas/actual/migracion/disponible", headers=headers)
    assert resp.json()["disponible"] is False

    resp = client.post(
        "/planillas/generar",
        json={
            "tipo": "quincenal",
            "periodo_inicio": "2026-09-01",
            "periodo_fin": "2026-09-15",
            "fecha_pago": "2026-09-16",
        },
        headers=headers,
    )
    assert resp.status_code == 201, resp.text
    planilla_id = resp.json()["id"]
    resp = client.get(f"/planillas/{planilla_id}/movimientos", headers=headers)
    assert resp.status_code == 200, resp.text
    movimientos = resp.json()
    assert len(movimientos) == 1
    assert movimientos[0]["salario_bruto"] == "450.00"


def test_migracion_rechaza_si_empresa_ya_tiene_datos(client, db):
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(filas_empleados=[_empleado_fila(identificacion, "2020-01-01", 700.00)])

    resp = _subir(client, headers, "/empresas/actual/migracion/confirmar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 422, resp.text


def test_migracion_rechaza_salario_bajo_el_minimo_historico(client, db):
    empresa = crear_empresa(db, _sufijo())
    usuario = crear_usuario(db, f"admin-{_sufijo()}@example.com")
    vincular(db, usuario, empresa, "admin")
    headers = login_y_seleccionar(client, usuario, empresa)
    crear_salario_minimo(db, datetime.date(2000, 1, 1), decimal.Decimal("800.00"), fecha_fin=None)

    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(filas_empleados=[_empleado_fila(identificacion, "2020-01-01", 100.00)])

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    errores = resp.json()["errores"]
    assert len(errores) == 1
    assert errores[0]["campo"] == "salario_base"


def test_migracion_rechaza_identificacion_duplicada_en_archivo(client, db):
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(
        filas_empleados=[
            _empleado_fila(identificacion, "2020-01-01", 700.00),
            _empleado_fila(identificacion, "2019-01-01", 750.00),
        ]
    )

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    errores = resp.json()["errores"]
    assert any("duplicada" in e["mensaje"] for e in errores)


def test_migracion_rechaza_identificacion_inexistente_en_hoja_opcional(client, db):
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(
        filas_empleados=[_empleado_fila(identificacion, "2020-01-01", 700.00)],
        filas_vacaciones=[{"identificacion": "no-existe", "dias_acumulados": 10, "dias_gozados": 0}],
    )

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    errores = resp.json()["errores"]
    assert any(e["hoja"] == migracion_service.HOJA_VACACIONES for e in errores)


def test_migracion_decimo_calcula_correcto_sin_hoja_2(client, db):
    """Ancla de seguridad #1: sin datos de 'cambios de salario recientes',
    el décimo del cuatrimestre en curso debe salir exacto usando el
    salario actual (no debe arrastrar ni un solo día de los 16 años de
    historia no migrados)."""
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(filas_empleados=[_empleado_fila(identificacion, "2010-03-15", 900.00)])

    resp = _subir(client, headers, "/empresas/actual/migracion/confirmar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text

    resp = client.get("/empleados", headers=headers)
    contrato_id = None
    for empleado in resp.json():
        if empleado["identificacion"] == identificacion:
            resp_c = client.get(f"/empleados/{empleado['id']}/contratos", headers=headers)
            contrato_id = resp_c.json()[0]["id"]
    assert contrato_id is not None

    # Cuatrimestre 'ago-dic' 2026 completo (16-ago a 15-dic) a $900/mes:
    # 120 días comerciales x $30/12 = $300.00 exacto (mismo caso de
    # verificación que ya usa CLAUDE.md para el motor de décimo).
    resp = client.post(
        "/planillas/generar-decimo",
        json={"cuatrimestre": "ago-dic", "anio": 2026, "fecha_pago": "2026-12-15"},
        headers=headers,
    )
    assert resp.status_code == 201, resp.text
    planilla_id = resp.json()["id"]
    resp = client.get(f"/planillas/{planilla_id}/movimientos", headers=headers)
    assert resp.status_code == 200, resp.text
    movimiento = resp.json()[0]
    assert movimiento["salario_bruto"] == "300.00"


def test_migracion_liquidacion_no_arrastra_anios_de_salario_pendiente(client, db):
    """Ancla de seguridad #2: liquidar a alguien poco después de migrado
    (sin haber corrido todavía ninguna planilla real) no debe contar
    'salario pendiente' desde el fecha_inicio real del contrato (2010)."""
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(filas_empleados=[_empleado_fila(identificacion, "2010-03-15", 900.00)])

    resp = _subir(client, headers, "/empresas/actual/migracion/confirmar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text

    resp = client.get("/empleados", headers=headers)
    contrato_id = None
    for empleado in resp.json():
        if empleado["identificacion"] == identificacion:
            resp_c = client.get(f"/empleados/{empleado['id']}/contratos", headers=headers)
            contrato_id = resp_c.json()[0]["id"]
    assert contrato_id is not None

    resp = client.post(
        f"/contratos/{contrato_id}/liquidacion",
        json={"motivo": "renuncia_voluntaria", "fecha_terminacion": "2026-09-10"},
        headers=headers,
    )
    assert resp.status_code == 201, resp.text
    liquidacion = resp.json()
    # 10 días desde el corte (1-sep) hasta el 10-sep a $30/día = $300.00 --
    # si la ancla de Planilla no existiera, esto arrastraría ~16 años.
    assert liquidacion["salario_pendiente"] == "300.00"


def test_migracion_con_cambio_de_salario_reciente(client, db):
    headers = _preparar_empresa_admin(db, client)
    fecha_corte = datetime.date(2026, 9, 1)
    identificacion = f"8-{_sufijo()}"
    contenido = _construir_excel(
        filas_empleados=[_empleado_fila(identificacion, "2015-01-01", 900.00)],
        filas_cambio_salario=[
            {
                "identificacion": identificacion, "salario_anterior": 800.00,
                "fecha_vigencia_desde": "2026-06-01", "fecha_vigencia_hasta": "2026-08-15",
            }
        ],
    )

    resp = _subir(client, headers, "/empresas/actual/migracion/validar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
    assert resp.json()["errores"] == []
    assert resp.json()["resumen"]["tramos_salario_adicionales"] == 1

    resp = _subir(client, headers, "/empresas/actual/migracion/confirmar", contenido, fecha_corte)
    assert resp.status_code == 200, resp.text
